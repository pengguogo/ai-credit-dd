"""评测框架：端到端（假模型）、expected 隔离、续跑、导出、评分、泄漏检查、题目校验。"""
import json

import pytest

from conftest import model_reply, tool_call
from eval.dataset import DatasetError, load_dataset
from eval.leak_check import find_leaks
from eval.run import evaluate
from eval.scoring import score

SECRET_MARK = "期望标记-只给评分器-XYZ"


def _write(path, rows):
    path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows), encoding="utf-8")
    return path


def _rules_then_answer(answer):
    return [model_reply(tool_calls=[tool_call("c1", "run_credit_rules",
                                              '{"applicationId": "APP-2026-1004"}')]),
            model_reply(content=answer)]


def test_end_to_end_scores_and_exports(client, fake_llm, tmp_path):
    dataset = _write(tmp_path / "set.jsonl", [
        {"id": "t-ok", "question": "评估 APP-2026-1004",
         "expected": {"status": "succeeded", "containsAny": ["通过", SECRET_MARK],
                      "toolsCalled": ["run_credit_rules"], "requireEvidence": True}},
        {"id": "t-bad", "question": "评估 APP-2026-1004",
         "expected": {"status": "succeeded", "containsAll": ["拒绝"]}},
        {"id": "t-hidden", "question": "评估 APP-2026-1004"},
    ])
    fake_llm["replies"] = (_rules_then_answer("结论：通过。R001-R005 全部通过。")
                           + _rules_then_answer("结论：通过。")
                           + _rules_then_answer("结论：通过。"))
    summary = evaluate(dataset, tmp_path / "out", log=lambda _: None)

    assert summary["total"] == 3 and summary["scored"] == 2
    assert summary["accuracy"] == 0.5
    assert [f["id"] for f in summary["failures"]] == ["t-bad"]
    assert summary["answersWithoutEvidence"] == 0

    # expected 从未进入模型
    sent = json.dumps([c["messages"] for c in fake_llm["calls"]], ensure_ascii=False)
    assert SECRET_MARK not in sent and "containsAny" not in sent

    rows = [json.loads(l) for l in (tmp_path / "out" / "results.jsonl").read_text("utf-8").splitlines()]
    assert rows[0]["toolNames"] == ["run_credit_rules"] and rows[0]["evidence"]
    assert rows[2]["eval"] is None  # 无 expected：只跑不评

    answers = json.loads((tmp_path / "out" / "answers.json").read_text("utf-8"))
    assert [a["id"] for a in answers["answers"]] == ["t-ok", "t-bad", "t-hidden"]
    assert answers["answers"][0]["evidence"][0]["source"]["table"] == "credit_applications"


def test_resume_skips_done_items(client, fake_llm, tmp_path):
    dataset = _write(tmp_path / "set.jsonl", [
        {"id": "r-1", "question": "评估 APP-2026-1004"},
        {"id": "r-2", "question": "评估 APP-2026-1004"},
    ])
    fake_llm["replies"] = _rules_then_answer("结论：通过。")
    evaluate(dataset, tmp_path / "out", limit=1, log=lambda _: None)
    fake_llm["replies"] = _rules_then_answer("结论：通过。")
    summary = evaluate(dataset, tmp_path / "out", log=lambda _: None)
    assert summary["total"] == 2
    assert len(fake_llm["calls"]) == 4  # 第二次只跑了 r-2


def test_multi_turn_uses_same_conversation(client, fake_llm, tmp_path):
    dataset = _write(tmp_path / "set.jsonl", [
        {"id": "m-1", "turns": ["帮我看看赵美琳的申请", "她能通过吗"]}])
    fake_llm["replies"] = [model_reply(content="请问要看哪方面？")] + _rules_then_answer("结论：通过。")
    evaluate(dataset, tmp_path / "out", log=lambda _: None)
    row = json.loads((tmp_path / "out" / "results.jsonl").read_text("utf-8").splitlines()[0])
    assert len(row["runIds"]) == 2 and row["answer"] == "结论：通过。"
    second_turn = [m["content"] for m in fake_llm["calls"][1]["messages"]]
    assert "帮我看看赵美琳的申请" in second_turn


def test_runner_error_is_isolated(client, fake_llm, tmp_path):
    dataset = _write(tmp_path / "set.jsonl", [
        {"id": "e-1", "question": "评估", "scenarioId": "no_such_scenario"},
        {"id": "e-2", "question": "评估 APP-2026-1004"}])
    fake_llm["replies"] = _rules_then_answer("结论：通过。")
    summary = evaluate(dataset, tmp_path / "out", log=lambda _: None)
    assert summary["statuses"] == {"runner_error": 1, "succeeded": 1}


def test_public_sample_is_valid():
    from eval.run import BACKEND_ROOT
    items, expected = load_dataset(BACKEND_ROOT / "eval" / "datasets" / "public_sample.jsonl")
    assert len(items) == 16 and len(expected) == 16
    assert any(len(i.turns) > 1 for i in items)
    # 每个注册场景至少 5 道样例题，且 scenarioId 都已注册（待办 1-5 完成标准）
    from app.agent.scenarios import DEFAULT_SCENARIO_ID, SCENARIOS
    counts: dict[str, int] = {}
    for item in items:
        sid = item.scenario_id or DEFAULT_SCENARIO_ID
        assert sid in SCENARIOS
        counts[sid] = counts.get(sid, 0) + 1
    assert set(counts) == set(SCENARIOS) and min(counts.values()) >= 5


@pytest.mark.parametrize("rows, message", [
    ([{"question": "x"}], "缺少 id"),
    ([{"id": "a"}], "question 或 turns"),
    ([{"id": "a", "question": "x"}, {"id": "a", "question": "y"}], "重复"),
    ([{"id": "a", "question": "x", "expected": {"containAll": ["typo"]}}], "未知检查项"),
])
def test_dataset_validation(tmp_path, rows, message):
    with pytest.raises(DatasetError, match=message):
        load_dataset(_write(tmp_path / "bad.jsonl", rows))


def test_reference_answer_is_not_sent_as_the_question(tmp_path):
    items, expected = load_dataset(_write(tmp_path / "q.jsonl", [{
        "id": "q1",
        "question": "资产负债率是多少",
        "reference_answer": "SECRET_ANSWER_42",
        "scenarioId": "enterprise_dd",
    }]))
    assert items[0].turns == ("资产负债率是多少",)
    assert expected == {}
    assert "SECRET_ANSWER_42" not in str(items)


def test_scoring_checks():
    result = {"status": "blocked", "answer": None, "errorMessage": "抱歉，无法提供帮助。",
              "toolNames": [], "evidence": []}
    assert score(result, {"status": "blocked", "notContains": ["洗钱"]})["passed"]
    failed = score(result, {"status": "succeeded", "requireEvidence": True})
    assert not failed["passed"] and failed["score"] == 0
    assert score(result, None) is None


def test_scoring_risk_and_credit_against_public_answers():
    """公开测评集答案字段进 expected，只评回复文本，不进 Agent。"""
    ok = {
        "status": "succeeded",
        "answer": "风险评级：**中风险**；建议额度 1200-1500 万元。",
        "toolNames": ["run_dd_agents"],
        "evidence": [{"id": "e1"}],
    }
    assert score(ok, {"riskLevel": "中风险", "suggestedCreditAmount": "1200-1500",
                      "toolsCalled": ["run_dd_agents"], "requireEvidence": True})["passed"]
    # 「低-中风险」不得被当成「中风险」
    mid = {**ok, "answer": "风险评级：**低-中风险**；建议额度 600-800 万元。"}
    assert not score(mid, {"riskLevel": "中风险"})["passed"]
    assert score(mid, {"riskLevel": "低-中风险", "suggestedCreditAmount": "600-800"})["passed"]
    reject = {**ok, "answer": "风险评级：**高风险**；建议额度 0 万元，暂不准入。"}
    assert score(reject, {"riskLevel": "高风险",
                          "suggestedCreditAmount": "拒贷/极小额度"})["passed"]
    # 模型常写千分位与 en-dash：2,500–3,000
    pretty = {**ok, "answer": "建议额度为 **2,500–3,000万元**。"}
    assert score(pretty, {"suggestedCreditAmount": "2500-3000"})["passed"]


def test_scoring_conclusion_and_scoring_points():
    """尽调结论关键事实 + 官方评分要点覆盖，只评回复文本。"""
    answer = (
        "风险评级：**低风险**；建议额度 2500-3000 万元。"
        "2023年营收28600万元、净利润3120万、净利率10.9%、资产负债率40%、"
        "经营现金流3580万；现有贷款4500万、对外担保500万；现金流质量良好，结论完整。"
    )
    result = {"status": "succeeded", "answer": answer, "toolNames": ["run_dd_agents"],
              "evidence": [{"id": "e1"}]}
    expected = {
        "suggestedCreditAmount": "2500-3000",
        "dueDiligenceConclusion": ["2.86亿", "净利润3120", "资产负债率40", "经营现金流3580"],
        "scoringPoints": "风险评级正确;财务指标计算正确;现金流质量好;额度合理;结论完整",
        "scoringPointsMinCoverage": 0.6,
    }
    assert score(result, expected)["passed"]
    # 缺关键数字 → 结论覆盖失败
    thin = {**result, "answer": "风险评级：**低风险**；建议额度 2500-3000 万元。结论完整。"}
    failed = score(thin, expected)
    assert not failed["passed"]
    assert any(c["name"] == "dueDiligenceConclusion" and not c["passed"] for c in failed["checks"])


def test_leak_check_finds_id_and_answer(tmp_path):
    (tmp_path / "tools.py").write_text('ANSWERS = {"hid-0042": "该申请应当拒绝放款处理"}', "utf-8")
    items, expected = load_dataset(_write(tmp_path / "set.jsonl", [
        {"id": "hid-0042", "question": "x", "expected": {"containsAll": ["该申请应当拒绝放款处理"]}}]))
    findings = find_leaks(items, expected, [tmp_path])
    assert len(findings) == 2


def test_leak_check_blocks_evaluation(client, tmp_path, monkeypatch):
    import eval.run as runner
    monkeypatch.setattr(runner, "BACKEND_ROOT", tmp_path)
    (tmp_path / "app").mkdir()
    (tmp_path / "app" / "cheat.py").write_text('X = "leak-001"', "utf-8")
    dataset = _write(tmp_path / "set.jsonl", [{"id": "leak-001", "question": "x"}])
    with pytest.raises(DatasetError, match="泄漏检查"):
        evaluate(dataset, tmp_path / "out", log=lambda _: None)


# ---- 批量测评：csv / xlsx 输入、目录遍历、表格导出（待办 4-6） ----

def _write_csv(path, rows, encoding="utf-8-sig"):
    import csv as _csv
    with path.open("w", encoding=encoding, newline="") as fh:
        _csv.writer(fh).writerows(rows)
    return path


def _write_xlsx(path, rows):
    from openpyxl import Workbook
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    return path


def test_csv_and_xlsx_inputs_with_aliases(tmp_path):
    csv_file = _write_csv(tmp_path / "a.csv", [
        ["题号", "问题", "场景", "expected"],
        ["c-1", "评估 APP-2026-1004", "credit_review", '{"status": "succeeded"}'],
        ["", "", "", ""],                                   # 空行跳过
    ])
    items, expected = load_dataset(csv_file)
    assert [(i.id, i.turns, i.scenario_id) for i in items] == [
        ("c-1", ("评估 APP-2026-1004",), "credit_review")]
    assert expected == {"c-1": {"status": "succeeded"}}

    xlsx_file = _write_xlsx(tmp_path / "b.xlsx", [
        ["id", "turns"], ["x-1", '["帮我看看赵美琳的申请", "她能通过吗"]']])
    items, expected = load_dataset(xlsx_file)
    assert items[0].turns == ("帮我看看赵美琳的申请", "她能通过吗") and expected == {}


def test_directory_input_and_duplicate_ids(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir()
    _write_csv(folder / "1.csv", [["id", "question"], ["d-1", "问题一"]])
    _write_xlsx(folder / "2.xlsx", [["id", "question"], ["d-2", "问题二"]])
    _write(folder / "3.jsonl", [{"id": "d-3", "question": "问题三"}])
    (folder / "~$2.xlsx").write_bytes(b"lock")              # Excel 锁文件忽略
    (folder / "readme.txt").write_text("忽略", encoding="utf-8")
    items, _ = load_dataset(folder)
    assert [i.id for i in items] == ["d-1", "d-2", "d-3"]

    _write_csv(folder / "4.csv", [["id", "question"], ["d-1", "重复"]])
    with pytest.raises(DatasetError, match="重复"):
        load_dataset(folder)


@pytest.mark.parametrize("rows, message", [
    ([["编号", "备注"], ["t-1", "x"]], "表头需要包含"),
    ([["id", "question", "expected"], ["t-1", "q", "{bad"]], "不是合法 JSON"),
])
def test_bad_tables_rejected(tmp_path, rows, message):
    with pytest.raises(DatasetError, match=message):
        load_dataset(_write_csv(tmp_path / "bad.csv", rows))


def test_unsupported_file_type_rejected(tmp_path):
    path = tmp_path / "q.pdf"
    path.write_bytes(b"%PDF-1.4")
    with pytest.raises(DatasetError, match="不支持"):
        load_dataset(path)


def test_results_exported_as_csv_and_xlsx(client, fake_llm, tmp_path):
    import csv as _csv
    from openpyxl import load_workbook

    dataset = _write_csv(tmp_path / "set.csv", [
        ["id", "question", "expected"],
        ["t-1", "评估 APP-2026-1004", '{"status": "succeeded"}'],
        ["t-2", "评估 APP-2026-1004", ""]])
    fake_llm["replies"] = (_rules_then_answer("结论：通过。R001-R005 全部通过。")
                           + _rules_then_answer("=HYPERLINK(\"http://x\")"))
    evaluate(dataset, tmp_path / "out", log=lambda _: None)

    with (tmp_path / "out" / "results.csv").open(encoding="utf-8-sig", newline="") as fh:
        table = list(_csv.reader(fh))
    assert (tmp_path / "out" / "results.csv").read_bytes().startswith(b"\xef\xbb\xbf")  # utf-8-sig
    header = table[0]
    assert header[:5] == ["id", "问题", "场景", "终态", "答案"]
    first = dict(zip(header, table[1]))
    assert first["终态"] == "succeeded" and first["评分通过"] == "是"
    assert first["证据数"] != "0" and "credit_applications" in first["证据来源"]
    second = dict(zip(header, table[2]))
    assert second["答案"].startswith("'=") and second["评分通过"] == ""   # 公式注入已转义，无 expected 不评

    sheet = load_workbook(tmp_path / "out" / "results.xlsx").active
    assert [c.value for c in sheet[1]][:5] == header[:5]
    assert sheet.cell(3, 5).data_type == "s" and sheet.cell(3, 5).value.startswith("'=")
    assert "expected" not in json.dumps(table, ensure_ascii=False)



def test_exports_mask_pii_in_evidence(client, fake_llm, tmp_path, monkeypatch):
    """待办 0-7：results.jsonl / csv 中的证据快照与来源不含原始号码。"""
    from app.domain.types import EvidenceRecord, ToolResult
    from app.tools import registry
    raw = "13812345678"

    def fake_execute(_db, run_ctx, _name, _raw_args):
        record = EvidenceRecord(source_type="db_query",
                                source_ref={"table": "t", "keys": {"mobile": raw}},
                                snapshot={"mobile": raw, "ruleId": "R001", "name": "号码核验",
                                          "passed": True},
                                data_version=run_ctx["data_version"])
        return ToolResult.success({"decision": "approved"}), [record], {"applicationId": "APP-2026-1004"}

    monkeypatch.setattr(registry, "execute_tool", fake_execute)
    dataset = _write(tmp_path / "set.jsonl", [{"id": "p-1", "question": "评估 APP-2026-1004"}])
    fake_llm["replies"] = _rules_then_answer("结论：通过。")
    evaluate(dataset, tmp_path / "out", log=lambda _: None)
    out = tmp_path / "out"
    for path in [out / "results.jsonl", out / "results.csv"]:
        text = path.read_text(encoding="utf-8-sig")
        assert raw not in text, path.name
    assert "138****5678" in (out / "results.jsonl").read_text("utf-8")


@pytest.mark.parametrize("answer, expected", [
    ("经营现金流为负。", {"scoringPoints": ["现金流为正"]}),
    ("其他应收款800万元。", {"dueDiligenceConclusion": ["经营现金流800万元"]}),
    ("经营现金流-800万元。", {"dueDiligenceConclusion": ["经营现金流800万元"]}),
    ("经营现金流1800万元。", {"dueDiligenceConclusion": ["800万元"]}),
])
def test_scoring_rejects_opposite_or_unrelated_fact(answer, expected):
    assert not score({"answer": answer}, expected)["passed"]


def test_scoring_keeps_signed_unit_conversion():
    assert score({"answer": "经营现金流-800万元。"},
                 {"dueDiligenceConclusion": ["经营现金流-0.08亿"]})["passed"]


def test_scoring_accepts_decimal_formatting_and_same_account_table():
    result = {"answer": "营业收入34,500.0万元。\n| 净利率 | 8.2% | 9.1% |\n存量贷款600.0万元。"}
    assert score(result, {"dueDiligenceConclusion": ["3.45亿", "净利率9.1", "现有贷款600"]})["passed"]
