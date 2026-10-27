"""题目加载与校验 — 题目与 expected 在这里分离。

@why 红线 R4/R5：expected 只交给评分器。加载后 Agent 侧只能拿到 Item（id、问题、场景），
     拿不到任何期望值；未知检查项直接报错，防止拼写错误让评分静默失效。
"""
from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path

KNOWN_CHECKS = {"status", "containsAll", "containsAny", "notContains",
                "toolsCalled", "requireEvidence",
                # 赛题 04 公开测评对照（只给评分器；不进泄漏扫描的 contains* 文本）
                "riskLevel", "suggestedCreditAmount",
                "dueDiligenceConclusion", "scoringPoints",
                "dueDiligenceConclusionMinCoverage", "scoringPointsMinCoverage"}

SUPPORTED_SUFFIXES = (".jsonl", ".json", ".csv", ".xlsx")

# 表格题目的列名别名（不区分大小写，去首尾空格）→ 内部字段。赛题数据包的表头不同时在这里补别名。
_COLUMN_ALIASES = {
    "id": "id", "编号": "id", "题号": "id", "序号": "id", "题目编号": "id",
    "question": "question", "问题": "question", "题目": "question", "问题描述": "question",
    "scenarioid": "scenarioId", "scenario": "scenarioId", "场景": "scenarioId",
    "turns": "turns", "多轮": "turns",
    "expected": "expected", "期望": "expected",
}


class DatasetError(ValueError):
    """题目文件格式错误。"""


@dataclass(frozen=True)
class Item:
    id: str
    turns: tuple[str, ...]          # 单轮题只有一个元素；多轮题按顺序在同一会话提问
    scenario_id: str | None = None


def load_dataset(path: str | Path) -> tuple[list[Item], dict[str, dict]]:
    """返回 (题目列表, {题目 ID: expected})。

    path 可以是单个文件，也可以是目录（按文件名顺序遍历其中所有支持的文件，题目 ID 全目录唯一）。
    支持 .jsonl、.json（列表或 {"items": [...]}）、.csv（utf-8 / utf-8-sig）、.xlsx（第一个工作表）。
    """
    path = Path(path)
    if not path.exists():
        raise DatasetError(f"题目文件不存在: {path}")
    files = _input_files(path)
    items: list[Item] = []
    expected: dict[str, dict] = {}
    seen: set[str] = set()
    for file in files:
        for index, row in enumerate(_read(file), 1):
            item, exp = _parse(row, f"{file.name} 第 {index} 题")
            if item.id in seen:
                raise DatasetError(f"题目 ID 重复: {item.id}")
            seen.add(item.id)
            items.append(item)
            if exp:
                expected[item.id] = exp
    if not items:
        raise DatasetError(f"题目文件为空: {path}")
    return items, expected


def _input_files(path: Path) -> list[Path]:
    if path.is_file():
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise DatasetError(f"不支持的题目文件类型 {path.suffix}，可用 {', '.join(SUPPORTED_SUFFIXES)}")
        return [path]
    # @why 跳过 Excel 打开文件时生成的 ~$ 锁文件
    files = sorted(p for p in path.iterdir()
                   if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES and not p.name.startswith("~$"))
    if not files:
        raise DatasetError(f"目录中没有可用的题目文件（{', '.join(SUPPORTED_SUFFIXES)}）: {path}")
    return files


def _read(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    text = path.read_text(encoding="utf-8")
    try:
        if path.suffix == ".jsonl":
            return [json.loads(line) for line in text.splitlines()
                    if line.strip() and not line.lstrip().startswith("//")]
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise DatasetError(f"{path.name} 不是合法 JSON: {exc}") from exc
    rows = data.get("items") if isinstance(data, dict) else data
    if not isinstance(rows, list):
        raise DatasetError(f"{path.name} 应为题目列表或 {{\"items\": [...]}}")
    return rows


def _read_csv(path: Path) -> list[dict]:
    # utf-8-sig 同时兼容带 BOM（Excel 另存）与不带 BOM 的文件
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        rows = [r for r in reader]
    return _table_rows(rows, path.name)


def _read_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = [["" if v is None else str(v) for v in r] for r in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _table_rows(rows, path.name)


def _table_rows(rows: list[list[str]], name: str) -> list[dict]:
    """表格 → 题目字典：首行为表头；turns / expected 列为 JSON 字符串；空行跳过。"""
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []
    header = [_COLUMN_ALIASES.get(str(h).strip().lower(), str(h).strip()) for h in rows[0]]
    if "id" not in header or not ({"question", "turns"} & set(header)):
        raise DatasetError(f"{name} 表头需要包含 id 与 question（或 turns）列，实际为 {rows[0]}")
    out = []
    for line_no, cells in enumerate(rows[1:], 2):
        record = {key: str(value).strip() for key, value in zip(header, cells) if str(value).strip()}
        for key in ("turns", "expected"):
            if key in record:
                try:
                    record[key] = json.loads(record[key])
                except json.JSONDecodeError as exc:
                    raise DatasetError(f"{name} 第 {line_no} 行 {key} 列不是合法 JSON: {exc}") from exc
        out.append(record)
    return out


def _parse(row: object, where: str) -> tuple[Item, dict | None]:
    if not isinstance(row, dict):
        raise DatasetError(f"{where}: 每道题必须是 JSON 对象")
    # 问询表自带的参考答案只给评分对照，不能进入送给 Agent 的题目
    row = {key: value for key, value in row.items()
           if str(key).strip().lower() not in {"reference_answer", "referenceanswer"}}
    item_id = str(row.get("id") or "").strip()
    if not item_id:
        raise DatasetError(f"{where}: 缺少 id")
    turns = row.get("turns")
    if turns is None:
        question = row.get("question")
        turns = [question] if question else []
    if not isinstance(turns, list) or not turns or not all(isinstance(t, str) and t.strip() for t in turns):
        raise DatasetError(f"{where}（{item_id}）: 需要非空的 question 或 turns")
    exp = row.get("expected")
    if exp is not None:
        if not isinstance(exp, dict):
            raise DatasetError(f"{where}（{item_id}）: expected 必须是对象")
        unknown = set(exp) - KNOWN_CHECKS
        if unknown:
            raise DatasetError(f"{where}（{item_id}）: 未知检查项 {sorted(unknown)}，可用 {sorted(KNOWN_CHECKS)}")
    return Item(id=item_id, turns=tuple(t.strip() for t in turns),
                scenario_id=row.get("scenarioId")), exp
