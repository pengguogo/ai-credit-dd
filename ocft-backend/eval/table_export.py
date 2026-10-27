"""结果表格导出 — results.csv（utf-8-sig）与 results.xlsx。

@why 《通用提交材料说明》1.1(2)：批量测评结果须以 excel / csv 输出。
     列定义集中在 COLUMNS；赛题数据包公布表头后，在这里改列名与取值，不动 runner。
     只导出题目、答案与可追溯信息，不导出 expected 与评分明细之外的内部字段。
"""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Callable

# (表头, 取值函数)。表头用中文便于评委直接阅读；runId 可在 /api/agent/runs/{runId} 追溯完整轨迹。
COLUMNS: list[tuple[str, Callable[[dict], object]]] = [
    ("id", lambda r: r["id"]),
    ("问题", lambda r: "\n".join(r.get("questions") or [])),
    ("场景", lambda r: r.get("scenarioId") or ""),
    ("终态", lambda r: r.get("status") or ""),
    ("答案", lambda r: r.get("answer") or r.get("errorMessage") or ""),
    ("错误码", lambda r: r.get("errorCode") or ""),
    ("调用工具", lambda r: ", ".join(r.get("toolNames") or [])),
    ("证据数", lambda r: len(r.get("evidence") or [])),
    ("证据来源", lambda r: "; ".join(_source(e) for e in r.get("evidence") or [])),
    ("runId", lambda r: ", ".join(r.get("runIds") or [])),
    ("耗时ms", lambda r: r.get("durationMs", 0)),
    ("评分通过", lambda r: "" if not r.get("eval") else ("是" if r["eval"]["passed"] else "否")),
]

# @why 防 CSV / Excel 公式注入：以 = + @ 开头，或 - 后不跟空格（区别于 Markdown 列表"- "）的文本，
#      在 Excel 中会被当成公式执行。前面加单引号，Excel 按文本显示。
_FORMULA_PREFIXES = ("=", "+", "@", "\t", "\r")


def _safe(value: object) -> object:
    if not isinstance(value, str) or not value:
        return value
    if value.startswith(_FORMULA_PREFIXES) or (value.startswith("-") and not value.startswith("- ")):
        return "'" + value
    return value


def _source(evidence: dict) -> str:
    source = evidence.get("source") or {}
    if not isinstance(source, dict):
        return str(source)
    target = source.get("table") or source.get("catalog") or evidence.get("sourceType") or ""
    keys = source.get("keys") or {}
    key_text = ",".join(f"{k}={v}" for k, v in keys.items() if v is not None)
    return f"{target}({key_text})" if key_text else str(target)


def table_rows(results: list[dict]) -> tuple[list[str], list[list[object]]]:
    header = [name for name, _ in COLUMNS]
    rows = [[_safe(getter(r)) for _, getter in COLUMNS] for r in results]
    return header, rows


def write_csv(results: list[dict], path: Path) -> None:
    header, rows = table_rows(results)
    # utf-8-sig：Excel 直接双击打开不乱码（提交材料说明第三章）
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        writer.writerows(rows)


def write_xlsx(results: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    header, rows = table_rows(results)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "results"
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                # @why openpyxl 会把 "=" 开头的字符串存成公式；显式标成文本
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {"问题": 40, "答案": 60, "证据来源": 40, "调用工具": 30, "runId": 20}
    for index, name in enumerate(header, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths.get(name, 14)
    sheet.freeze_panes = "A2"
    workbook.save(path)
